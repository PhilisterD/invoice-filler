import os
import re
import logging
from copy import deepcopy
from datetime import datetime

import pandas as pd
from docx.oxml import OxmlElement
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from utils import num_to_chinese, format_date


logger = logging.getLogger('invoice-filler')


def is_yellow_highlight(run) -> bool:
    """判断一个 run 是否是黄色高亮 (highlight_color == 7)"""
    if run.font.highlight_color == 7:
        return True
    rPr = run._element.find(qn('w:rPr'))
    if rPr is not None:
        shd = rPr.find(qn('w:shd'))
        if shd is not None:
            fill = shd.get(qn('w:fill'))
            if fill and fill.upper() in ('FFFF00', 'FF0', 'YELLOW'):
                return True
    return False


def find_yellow_runs(paragraph):
    """返回段落中所有黄色高亮的 run 索引"""
    return [i for i, run in enumerate(paragraph.runs) if is_yellow_highlight(run)]


def find_table_rows_with_yellow(table):
    """找出表格中包含黄色高亮的行索引"""
    rows = []
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if is_yellow_highlight(run):
                        rows.append(ri)
                        break
                if ri in rows:
                    break
            if ri in rows:
                break
    return sorted(set(rows))


SUMMARY_TEXT_HINTS = ('大写', '人民币', '总计', '合计', '费用')


def get_row_text(row) -> str:
    """合并行内所有单元格文本，用于结构识别。"""
    texts = []
    for cell in row.cells:
        text = cell.text.strip()
        if text and text not in texts:
            texts.append(text)
    return ' '.join(texts).strip()


def is_summary_row_text(text: str) -> bool:
    """根据行文本判断是否像汇总行。"""
    if not text:
        return False
    return any(hint in text for hint in SUMMARY_TEXT_HINTS)


def extract_summary_label(summary_row) -> str:
    """从模板汇总行里提取前缀文案，便于适配不同模板。"""
    raw_text = get_row_text(summary_row)
    if not raw_text:
        return '费用总计'

    label = raw_text
    label = re.sub(r'人民币\s*[\d,]+(?:\.\d+)?\s*元?', '', label)
    label = re.sub(r'（?\s*大写[:：]?\s*人民币[^）]*）?', '', label)
    label = re.sub(r'[\d,]+(?:\.\d+)?', '', label)
    label = re.sub(r'\s+', '', label)
    label = label.strip('：:，,。 ')
    return label or '费用总计'


def build_summary_text(total_sum, summary_row=None) -> str:
    """生成通用汇总文案，保留模板原有前缀。"""
    label = extract_summary_label(summary_row) if summary_row is not None else '费用总计'
    label = label.rstrip('：:，,。 ')
    chinese_total = num_to_chinese(int(total_sum))
    return f'{label}：人民币  {int(total_sum)}元（大写：人民币{chinese_total}）'


def detect_table_layout(table):
    """识别数据行与汇总行。"""
    yellow_rows = find_table_rows_with_yellow(table)
    if len(yellow_rows) >= 2:
        return yellow_rows[0], yellow_rows[-1], yellow_rows
    if len(table.rows) >= 3:
        return 1, len(table.rows) - 1, yellow_rows
    return None, None, yellow_rows


def can_access_grid_offset(table, row_idx, col_idx) -> bool:
    """判断表格某行是否真的存在指定网格列。"""
    try:
        table.rows[row_idx]._tr.tc_at_grid_offset(col_idx)
        return True
    except Exception:
        return False


def set_vertical_merge(cell, value):
    """设置单元格纵向合并标记。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    v_merge = tc_pr.find(qn('w:vMerge'))
    if v_merge is None:
        v_merge = OxmlElement('w:vMerge')
        tc_pr.append(v_merge)
    if value is None:
        if v_merge.get(qn('w:val')) is not None:
            del v_merge.attrib[qn('w:val')]
    else:
        v_merge.set(qn('w:val'), value)


def merge_vertical_group(table, row_indices, col_idx):
    """对连续行分组执行纵向合并。"""
    row_indices = sorted(set(row_indices))
    if len(row_indices) < 2:
        return False

    merged = False
    group_start = row_indices[0]
    group_end = row_indices[0]

    def close_group(start_idx, end_idx):
        nonlocal merged
        if start_idx >= end_idx:
            return
        if col_idx >= len(table.rows[start_idx].cells) or col_idx >= len(table.rows[end_idx].cells):
            return
        first_cell = table.rows[start_idx].cells[col_idx]
        set_vertical_merge(first_cell, 'restart')
        for ri in range(start_idx + 1, end_idx + 1):
            if col_idx >= len(table.rows[ri].cells):
                continue
            set_vertical_merge(table.rows[ri].cells[col_idx], 'continue')
        merged = True

    for row_idx in row_indices[1:]:
        if row_idx == group_end + 1:
            group_end = row_idx
            continue
        close_group(group_start, group_end)
        group_start = group_end = row_idx

    close_group(group_start, group_end)
    return merged


def merge_total_column_span(table, row_indices, col_idx) -> bool:
    """安全合并总计列：只合并有数值合计的明细行。"""
    row_indices = sorted(set(row_indices))
    if len(row_indices) < 2:
        return False

    if merge_vertical_group(table, row_indices, col_idx):
        return True

    candidate_spans = []
    if len(row_indices) >= 3:
        candidate_spans.append((row_indices[1], row_indices[-2]))
    candidate_spans.append((row_indices[0], row_indices[-1]))

    for span_start, span_end in candidate_spans:
        if span_start >= span_end:
            continue

        if col_idx >= len(table.rows[span_start].cells) or col_idx >= len(table.rows[span_end].cells):
            continue

        first_total_cell = table.rows[span_start].cells[col_idx]
        last_total_cell = table.rows[span_end].cells[col_idx]
        if first_total_cell == last_total_cell:
            return True

        try:
            first_total_cell.merge(last_total_cell)
            return True
        except ValueError:
            continue

    return False


def has_numeric_value(value) -> bool:
    """判断值是否为可用于合并依据的数值。"""
    if value is None:
        return False
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return False
    try:
        return not pd.isna(value)
    except Exception:
        return False


def set_cell_text(cell, text):
    """安全设置单元格文本，保留第一个 run 的样式"""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def remove_highlight_from_run(run):
    """清除 run 的高亮和底纹"""
    run.font.highlight_color = None
    rPr = run._element.find(qn('w:rPr'))
    if rPr is not None:
        for hl in rPr.findall(qn('w:highlight')):
            rPr.remove(hl)
        for shd in rPr.findall(qn('w:shd')):
            rPr.remove(shd)


def remove_all_highlights(doc):
    """清除文档中所有段落和表格的高亮"""
    for p in doc.paragraphs:
        for run in p.runs:
            remove_highlight_from_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        remove_highlight_from_run(run)


def copy_row_xml(table, source_row_idx):
    """复制表格中的一行并插入到其后（XML层面）。"""
    tbl = table._tbl
    source_tr = table.rows[source_row_idx]._tr
    new_tr = deepcopy(source_tr)
    actual_idx = list(tbl).index(source_tr)
    tbl.insert(actual_idx + 1, new_tr)


def fill_template(excel_path: str, template_path: str, output_dir: str, logger=None) -> list[dict]:
    """
    读取Excel，按受托方分组，为每组生成一个Word文档。
    返回生成文件的信息列表。
    """
    active_logger = logger or globals()['logger']

    active_logger.info('读取 Excel: %s', excel_path)
    df = pd.read_excel(excel_path)
    rename_map = {}
    for col in df.columns:
        c = str(col).strip()
        if '受托方' in c or '委托方' in c:
            rename_map[col] = '受托方'
        elif '姓名' in c:
            rename_map[col] = '姓名'
        elif '服务内容' in c:
            rename_map[col] = '服务内容'
        elif '数量' in c and '单价' not in c and '合计' not in c:
            rename_map[col] = '数量'
        elif '单价' in c:
            rename_map[col] = '单价'
        elif '合计' in c and '总计' not in c:
            rename_map[col] = '合计'
        elif '总计' in c:
            rename_map[col] = '总计'
        elif '日期' in c:
            rename_map[col] = '日期'
        elif '序号' in c:
            rename_map[col] = '序号'
    df.rename(columns=rename_map, inplace=True)

    active_logger.info('识别到列: %s', list(df.columns))

    for col in ['数量', '单价', '合计', '总计']:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce')

    grouped = df.groupby('受托方', sort=False)
    results = []

    active_logger.info('开始分组处理: %s 个受托方', grouped.ngroups)

    for party, group in grouped:
        try:
            group = group.sort_values('序号') if '序号' in group.columns else group
            active_logger.info('处理受托方: %s, 记录数=%s', party, len(group))
            doc = Document(template_path)

            # ---- 1. 替换文档中黄色高亮的文本（委托方、日期） ----
            active_logger.info('替换模板高亮文本')
            for p in doc.paragraphs:
                yellow_indices = find_yellow_runs(p)
                for idx in yellow_indices:
                    run = p.runs[idx]
                    text = run.text
                    if any(kw in text for kw in ['研究所', '大学', '公司', '学院', '中心']):
                        run.text = str(party)
                    elif re.search(r'\d{4}年\d{1,2}月\d{1,2}日', text):
                        date_val = group['日期'].iloc[0]
                        run.text = format_date(date_val)

            # ---- 2. 处理表格 ----
            if not doc.tables:
                active_logger.warning('模板中没有表格: %s', template_path)
                results.append({
                    'party': party,
                    'filename': f"{party}_明细.docx",
                    'path': None,
                    'error': '模板中没有表格'
                })
                continue

            table = doc.tables[0]
            data_row_idx, summary_row_idx, yellow_rows = detect_table_layout(table)
            active_logger.info('表格布局识别: data_row_idx=%s summary_row_idx=%s yellow_rows=%s',
                               data_row_idx, summary_row_idx, yellow_rows)

            if data_row_idx is None or summary_row_idx is None:
                active_logger.error('无法识别模板中的数据行和汇总行')
                results.append({
                    'party': party,
                    'filename': f"{party}_明细.docx",
                    'path': None,
                    'error': '无法识别模板中的数据行和汇总行'
                })
                continue

            # 获取列结构
            header_row = table.rows[0]
            headers = [cell.text.strip() for cell in header_row.cells]
            col_map = {}
            for i, h in enumerate(headers):
                if '服务' in h or '内容' in h:
                    col_map['服务内容'] = i
                elif '数量' in h:
                    col_map['数量'] = i
                elif '单价' in h:
                    col_map['单价'] = i
                elif '合计' in h and '总计' not in h:
                    col_map['合计'] = i
                elif '总计' in h:
                    col_map['总计'] = i

            records = group.to_dict('records')
            active_logger.info('列映射: %s', col_map)

            # 删除多余数据行，只保留第一个作为模板
            existing_data_rows = max(1, summary_row_idx - data_row_idx)
            tbl = table._tbl
            for _ in range(existing_data_rows - 1):
                # 删除 data_row_idx + 1（第二个数据行），反复删直到只剩一个
                tr_to_remove = table.rows[data_row_idx + 1]._tr
                tbl.remove(tr_to_remove)

            # 清空保留的模板行
            data_row = table.rows[data_row_idx]
            for cell in data_row.cells:
                set_cell_text(cell, "")

            # 插入副本（需要记录数 - 1 个新行）
            for _ in range(len(records) - 1):
                copy_row_xml(table, data_row_idx)

            # 重新打开文档以刷新 table.rows（python-docx 缓存问题）
            temp_path = os.path.join(output_dir, '_temp.docx')
            doc.save(temp_path)
            doc = Document(temp_path)
            os.remove(temp_path)
            table = doc.tables[0]

            # 重新定位行索引
            data_row_idx, summary_row_idx, _ = detect_table_layout(table)
            if data_row_idx is None or summary_row_idx is None:
                active_logger.warning('刷新后无法重新识别数据行或汇总行: %s', party)
                continue

            # 填充数据行
            total_sum = 0
            merge_row_indices = []
            for i, rec in enumerate(records):
                target_row = table.rows[data_row_idx + i]
                service = str(rec.get('服务内容', ''))
                qty = rec.get('数量', 0)
                price = rec.get('单价', 0)
                subtotal = rec.get('合计', 0)
                if pd.isna(subtotal) or subtotal == 0:
                    subtotal = (qty or 0) * (price or 0)
                total_sum += subtotal
                active_logger.info('写入明细行 %s/%s: 服务=%s 数量=%s 单价=%s 合计=%s',
                                   i + 1, len(records), service, qty, price, subtotal)
                if has_numeric_value(subtotal):
                    merge_row_indices.append(data_row_idx + i)

                cells = target_row.cells
                if '服务内容' in col_map and col_map['服务内容'] < len(cells):
                    set_cell_text(cells[col_map['服务内容']], service)
                if '数量' in col_map and col_map['数量'] < len(cells):
                    set_cell_text(cells[col_map['数量']], str(int(qty)) if qty == int(qty) else str(qty))
                if '单价' in col_map and col_map['单价'] < len(cells):
                    set_cell_text(cells[col_map['单价']], str(int(price)) if price == int(price) else str(price))
                if '合计' in col_map and col_map['合计'] < len(cells):
                    set_cell_text(cells[col_map['合计']], str(int(subtotal)) if subtotal == int(subtotal) else str(subtotal))

            active_logger.info('明细写入完成: total_sum=%s', total_sum)

            # 填充总计列（只在第一行设置，然后安全地尝试纵向合并）
            total_col = col_map.get('总计', len(table.rows[0].cells) - 1)
            for ri in range(data_row_idx, summary_row_idx):
                row = table.rows[ri]
                if total_col < len(row.cells):
                    cell = row.cells[total_col]
                    if ri == data_row_idx:
                        set_cell_text(cell, str(int(total_sum)) if total_sum == int(total_sum) else str(total_sum))
                    else:
                        set_cell_text(cell, "")

            # 纵向合并总计列单元格：只合并“合计/元”有数值的明细行
            merge_total_column_span(table, merge_row_indices, total_col)

            # 填充汇总行（合并单元格中多个 cell 可能共享同一个 XML 元素，需要去重）
            summary_row = table.rows[summary_row_idx]
            if summary_row.cells:
                seen_tc = set()
                summary_text = build_summary_text(total_sum, summary_row)
                for cell in summary_row.cells:
                    tc = cell._tc
                    if tc in seen_tc:
                        continue
                    seen_tc.add(tc)
                    set_cell_text(cell, summary_text)

            # ---- 4. 清除所有黄色高亮 ----
            remove_all_highlights(doc)

            # ---- 5. 保存 ----
            person = str(group['姓名'].iloc[0]) if '姓名' in group.columns else ''
            safe_person = re.sub(r'[\n\r\\/:*?"<>|]', '_', person).strip()
            safe_party = re.sub(r'[\n\r\\/:*?"<>|]', '_', str(party)).strip()
            filename = f"{safe_person}_{safe_party}_{int(total_sum)}元明细.docx"
            out_path = os.path.join(output_dir, filename)
            doc.save(out_path)
            active_logger.info('文件保存完成: %s', out_path)
            results.append({
                'party': party,
                'filename': filename,
                'path': out_path,
                'records': len(records),
                'total': int(total_sum)
            })
        except Exception:
            active_logger.exception('处理受托方失败: %s', party)
            results.append({
                'party': party,
                'filename': f"{party}_明细.docx",
                'path': None,
                'error': '生成过程中发生异常，请查看日志'
            })

    return results


if __name__ == '__main__':
    import sys
    if len(sys.argv) >= 3:
        excel = sys.argv[1]
        template = sys.argv[2]
        out = sys.argv[3] if len(sys.argv) > 3 else './output'
    else:
        excel = '../明细-excel.xlsx'
        template = '../早鸟天筹明细-模板1.docx'
        out = './output'
    os.makedirs(out, exist_ok=True)
    res = fill_template(excel, template, out)
    for r in res:
        print(r)
